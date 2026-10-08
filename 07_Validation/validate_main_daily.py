"""Final Output Sales tab: date range (calendar Date From / Date To inside the September Comparison Table) + Download Excel.

The separate "Daily comparison (data snapshots)" section was removed (2026-10-07); this checks it is gone.
Expected values are recomputed independently:
  * eBay IDs / filters: dataset.json (validate_html.expected_ids) + ph_map + PH Sheet categories (build.sheet_categories)
  * full month 1-30 Sep: dataset.json monthly figures (exact reconciliation)
  * any range: a direct read-only query of order_management_copy (DATABASE_URL) for exactly those days:
      TY = the selected days of 2026 (1 Sep .. latest order date, e.g. 7 Oct), LY = the same calendar days of 2025
      Sales = SUM(order_total), Orders = COUNT(DISTINCT order_id) of Completed eBay orders (order_transaction, order_date)
      Views = SUM(click) (traffic_data, date); Conversion % = Orders / Views x 100, Views 0 -> N/A
Writes reports/main_daily_validation.json and reports/date_range_validation_<today>.json; exit 1 on any failure.
"""
import datetime, json, os, pathlib, re, sys, tempfile
from decimal import Decimal, ROUND_HALF_UP
import openpyxl, psycopg2
from playwright.sync_api import sync_playwright
from validate_html import HTML, expected_ids

BASE = pathlib.Path(__file__).resolve().parent.parent
REP = BASE / "06_Validation" / "reports"
PH_MAP = BASE / "04_HTML_Report" / "data" / "ph_map.json"
sys.path.insert(0, str(BASE / "04_HTML_Report"))
from build import sheet_categories      # noqa: E402  (same PH Sheet category source as the page)
NOPH = "__none__"
SS = ("so_926407", "led_sone", "electricalsone", "huettenlampen", "ledsonede")
DAILY = BASE / "04_HTML_Report" / "data" / "daily_detail.json"
_dd = json.load(open(DAILY, encoding="utf-8"))
LAST, END, TRF_END = _dd["last_day"], _dd["end_date"], _dd["traffic_end"]; del _dd
# day index i = days since 1 Sep + 1 (1-30 Sep, 31.. Oct); (1, LAST) first: its traffic presence decides Views null-ness past Sep
RANGES = [(1, LAST), (1, 30), (1, 15), (16, 30), (7, 7)] + ([(31, LAST), (16, LAST), (LAST, LAST)] if LAST > 30 else [])
TRF_LAST = (datetime.date.fromisoformat(TRF_END) - datetime.date(2026, 9, 1)).days + 1     # day index of the latest traffic date
results, figures = [], {}


def check(name, ok, detail=""):
    results.append({"check": name, "ok": bool(ok), "detail": str(detail)[:800]})
    print(("PASS " if ok else "FAIL ") + name + ("" if ok else f"  -> {detail}"))


def iso(i, y=2026):
    return (datetime.date(y, 9, 1) + datetime.timedelta(days=i - 1)).isoformat()


def lbl(a, b, y):
    dm = lambda i: (i, "Sep") if i <= 30 else (i - 30, "Oct")
    (x, m), (z, n) = dm(a), dm(b)
    return (f"{x} {m}" if a == b else f"{x}–{z} {m}" if m == n else f"{x} {m}–{z} {n}") + f" {y}"


def db_range(cur, a, b):
    """{(mp, item_id, 'LY'|'TY'): [sales, orders, views or None]} for day index a..b in 2025 and in 2026."""
    w = lambda c: (f"(({c} >= '{iso(a, 2025)}' AND {c} < DATE '{iso(b, 2025)}' + 1) OR "
                   f"({c} >= '{iso(a)}' AND {c} < DATE '{iso(b)}' + 1))")
    out = {}
    cur.execute(f"""SELECT market_place, item_id, CASE WHEN order_date < '2026-01-01' THEN 'LY' ELSE 'TY' END,
                           SUM(COALESCE(order_total,0))::float, COUNT(DISTINCT order_id)
                    FROM public.order_transaction WHERE source_name='EBAY' AND order_status='Completed'
                    AND ss_name IN %s AND {w('order_date')} GROUP BY 1,2,3""", (SS,))
    for mp, i, p, s, o in cur.fetchall():
        out[(mp, i, p)] = [s, o, None]
    cur.execute(f"""SELECT market_place, ref_id, CASE WHEN date < '2026-01-01' THEN 'LY' ELSE 'TY' END, SUM(COALESCE(click,0))::float
                    FROM public.traffic_data WHERE which_channel=2 AND sub_source_name IN %s AND {w('date')} GROUP BY 1,2,3""", (SS,))
    for mp, i, p, v in cur.fetchall():
        out.setdefault((mp, i, p), [0.0, 0, None])[2] = v
    return out


def main():
    ids = expected_ids(); phm = json.load(open(PH_MAP, encoding="utf-8"))["map"]; cats = sheet_categories()
    pool = [dict(d, ph=phm.get(d["id"], []), cat=cats.get(d["id"], "")) for d in ids.values()]
    by_key = {(d["id"], d["mp"]): d for d in pool}

    def filt(ph="", acct="", seg="", q="", cat=""):
        q = q.lower()
        ok_ph = (lambda d: bool(d["ph"])) if ph == "" else (lambda d: not d["ph"]) if ph == NOPH else (lambda d: ph in d["ph"])
        return [d for d in pool if ok_ph(d) and (not acct or d["acct"] == acct) and (not seg or seg in d["seg"]) and (not cat or d["cat"] == cat)
                and (not q or q in d["id"].lower() or any(q in s.lower() for s in list(d["skus"]) + list(d["lsku"] or [])))]

    con = psycopg2.connect(os.environ["DATABASE_URL"]); con.set_session(readonly=True, autocommit=True); cur = con.cursor()
    cur.execute("select current_database()"); dbname = cur.fetchone()[0]
    src = {r: db_range(cur, *r) for r in RANGES}
    con.close()
    check("source DB for the day ranges is order_management_copy (same as extract.py / extract_daily.py)", dbname == "order_management_copy", dbname)

    def expect(d, rng, p):
        """[sales, orders, views] the page must show for this eBay ID, period and range (views None = no traffic row in the month)."""
        if rng == (1, 30):
            x = d[p]; return [x["s"], x["o"], x["v"]]
        s, o, v = src[rng].get((d["mp"], d["id"], p), [0.0, 0, None])
        if rng[1] <= 30:                       # inside September: the monthly "no traffic row" rule
            return [s, o, None if d[p]["v"] is None else (v or 0)]
        any_trf = src[(1, LAST)].get((d["mp"], d["id"], p), [0, 0, None])[2] is not None
        if p == "TY" and rng[0] > TRF_LAST:    # no 2026 traffic at all for these days yet: No data, never 0
            return [s, o, None]
        return [s, o, (v or 0) if any_trf else None]

    with sync_playwright() as pw:
        br = pw.chromium.launch(); pg = br.new_page(viewport={"width": 1600, "height": 1000}, accept_downloads=True)
        errors = []; pg.on("console", lambda m: errors.append(m.text) if m.type == "error" else None); pg.on("pageerror", lambda e: errors.append(str(e)))
        pg.goto(HTML.as_uri()); pg.wait_for_selector("#t tbody tr[data-id]")

        # ---- structure ----
        main_txt = pg.inner_text("#tabMain")
        check("Daily comparison (data snapshots) section removed (no card, table, checkbox, snapshot selects or wording)",
              pg.locator("#dCard, #dt, #dOnly, #dFrom, #dTo, #dNote, #dCount").count() == 0
              and not re.search(r"Daily comparison|[Ss]napshots stored|Only eBay IDs that changed", main_txt), "")
        info = pg.evaluate("""()=>{const f=document.querySelector('#sdFrom'),t=document.querySelector('#sdTo'),c=f.closest('.card');
            return {types:[f.type,t.type],min:[f.min,t.min],max:[f.max,t.max],val:[f.value,t.value],
                    card:c?c.querySelector('h2').innerText:null,sameCard:c===document.querySelector('#t').closest('.card'),
                    inFilterBar:!!f.closest('.filters'),tables:document.querySelectorAll('#tabMain table').length}}""")
        check(f"Date From / Date To are native calendar inputs (type=date) from 1 Sep 2026 to the latest order date {END} (traffic to {TRF_END})",
              info["types"] == ["date", "date"] and info["min"] == ["2026-09-01"] * 2 and info["max"] == [END] * 2, info)
        check("date controls sit inside the September Comparison Table card; no second table on the tab",
              info["card"] and info["card"].upper() == "SEPTEMBER COMPARISON TABLE" and info["sameCard"] and not info["inFilterBar"] and info["tables"] == 1, info)
        check(f"default range = 1 Sep 2026 → {END} (latest data date), header and count label show it", info["val"] == ["2026-09-01", END]
              and f"LY = {lbl(1, LAST, 2025)} · TY = {lbl(1, LAST, 2026)}" in pg.inner_text("#sub") and "rng" in pg.inner_html("#count"), (info["val"], pg.inner_text("#sub")))
        check("Price & SKU tab has no sales date-range or daily comparison controls", pg.locator("#tabPrice #sdFrom, #tabPrice #fDl, #tabPrice #dt").count() == 0)

        def main_n():
            return int(re.search(r"of ([\d,]+) eBay IDs", pg.inner_text("#count")).group(1).replace(",", ""))

        def download():
            with pg.expect_download() as dl: pg.click("#fDl")
            p = pathlib.Path(tempfile.gettempdir()) / dl.value.suggested_filename; dl.value.save_as(p)
            return dl.value.suggested_filename, openpyxl.load_workbook(p)

        def set_range(a, b):
            pg.fill("#sdFrom", iso(a)); pg.fill("#sdTo", iso(b))
            # set again in order so From <= To holds whichever input changed last
            pg.fill("#sdFrom", iso(a)); pg.wait_for_timeout(500)

        def set_filters(ph="", acct="", seg="", q="", cat=""):
            pg.select_option("#fPh", ph); pg.select_option("#fAcct", acct); pg.select_option("#fSeg", seg); pg.select_option("#fCat", cat)
            pg.fill("#fQ", q); pg.wait_for_timeout(450)

        def close(g, v, tol):
            return (g is None) == (v is None) and (g is None or abs(g - v) <= tol)

        def verify_excel(wb, fname, sel, rng):
            a, b = rng; bad = []
            if wb.sheetnames != ["Final Output Sales"]: bad.append(("sheets", wb.sheetnames))
            rows = list(wb["Final Output Sales"].iter_rows(values_only=True)); H = list(rows[0]); ix = {h: i for i, h in enumerate(H)}
            col = {(p, m): f"{p} {m} ({lbl(a, b, 2025 if p == 'LY' else 2026)})" for p in ("LY", "TY") for m in ("Sales", "Orders", "Views", "Conversion %")}
            if any(c not in ix for c in col.values()) or f"YoY % ({lbl(a, b, 2026)} vs {lbl(a, b, 2025)})" not in ix:
                return [("headers", H[7:20])], {}
            if f"_{iso(a)}_to_{iso(b)}.xlsx" not in fname: bad.append(("filename", fname))
            if len(rows) - 1 != len(sel): bad.append(("rows", len(sel), len(rows) - 1))
            want = {(d["id"], d["mp"]): d for d in sel}; got = {}
            for r in rows[1:]:
                k = (r[0], r[ix["Marketplace"]]); d = want.get(k)
                if not d: bad.append(("unexpected row", k)); continue
                got[k] = {}
                for p in ("LY", "TY"):
                    es, eo, ev = expect(d, rng, p)
                    gs, go_, gv, gc = (r[ix[col[(p, m)]]] for m in ("Sales", "Orders", "Views", "Conversion %"))
                    got[k][p] = (gs, go_, gv)
                    if p == "LY" and d["nl"]:
                        if gv != "Not live in Sep 2025": bad.append((k, p, "not-live views", gv))
                        if not close(gs, round(es, 2), 0.0051) or go_ != eo: bad.append((k, p, "sales/orders", (round(es, 2), eo), (gs, go_)))
                        continue
                    if not close(gs, round(es, 2), 0.0051) or go_ != eo: bad.append((k, p, "sales/orders", (round(es, 2), eo), (gs, go_)))
                    if not close(gv, ev, 0.5): bad.append((k, p, "views", ev, gv))
                    # half-up like the page (1/32 x 100 = 3.125 -> 3.13); Python round() would give banker's 3.12
                    ec = None if ev is None else "N/A" if ev == 0 else float((Decimal(eo) * 100 / Decimal(repr(ev))).quantize(Decimal("0.01"), ROUND_HALF_UP))
                    if (gc != ec) if isinstance(ec, str) or ec is None else not close(gc, ec, 0.0051): bad.append((k, p, "conversion", ec, gc))
            return bad, got

        # ---- every eBay ID (All PHs + Unassigned = all 14,912), every range, against the DB / monthly extract ----
        excel = {}
        for rng in RANGES:
            set_range(*rng)
            a, b = rng; full = rng == (1, 30)
            hd = pg.eval_on_selector_all("#t thead th", "e=>e.map(x=>x.firstChild.textContent+(x.querySelector('.rh')?'|'+x.querySelector('.rh').textContent:''))")
            lab_ok = (all("|" not in h for h in hd) if full else
                      f"LY Sales|{lbl(a, b, 2025)}" in hd and f"TY Sales|{lbl(a, b, 2026)}" in hd and f"TY Conversion %|{lbl(a, b, 2026)}" in hd
                      and not any("|" in h for h in hd if h.split("|")[0] in ("Ad Impressions", "Ad Sales", "Segment", "LY Ad Sales", "TY Ad Sales")))
            sub_ok = f"LY = {lbl(a, b, 2025)} · TY = {lbl(a, b, 2026)}" in pg.inner_text("#sub")
            bad, tot = [], {"LY": [0.0, 0, 0], "TY": [0.0, 0, 0]}
            for ph in ("", NOPH):
                set_filters(ph=ph); sel = filt(ph=ph)
                if main_n() != len(sel): bad.append((ph or "All PHs", "table rows", len(sel), main_n()))
                fname, wb = download(); b_, got = verify_excel(wb, fname, sel, rng); bad += b_
                for k, v in got.items():
                    excel.setdefault(rng, {})[k] = v
                    for p in ("LY", "TY"):
                        tot[p][0] += v[p][0] or 0; tot[p][1] += v[p][1] or 0; tot[p][2] += v[p][2] if isinstance(v[p][2], (int, float)) else 0
            # IDs with 2 PHs appear once under All PHs; All PHs + Unassigned = every eBay ID exactly once
            n_all = len(excel.get(rng, {}))
            figures[lbl(a, b, 2026)] = {"eBay IDs": n_all, **{f"{p} {m}": (round(tot[p][i], 2) if i == 0 else tot[p][i]) for p in ("LY", "TY") for i, m in enumerate(("Sales", "Orders", "Views"))}}
            # rendered table, first page: TY / LY Sales cells equal the expected range values
            set_filters(); rend = pg.eval_on_selector_all("#t tbody tr[data-id]", "rs=>rs.map(r=>[r.dataset.id,r.dataset.mp,r.querySelector('td[data-k=lys]').innerText,r.querySelector('td[data-k=tys]').innerText,r.querySelector('td[data-k=tyo]').innerText])")
            for i, mp, ls, ts, to in rend:
                d = by_key[(i, mp)]
                nm = lambda t: float(t.replace(",", ""))
                if abs(nm(ts) - expect(d, rng, "TY")[0]) > 0.0051 or abs(nm(ls) - expect(d, rng, "LY")[0]) > 0.0051 or nm(to) != expect(d, rng, "TY")[1]:
                    bad.append(("rendered", i, mp, ls, ts, to))
            txt = pg.inner_text("#t")
            check(f"{lbl(a, b, 2026)} vs {lbl(a, b, 2025)}: every eBay ID ({n_all}) — table rows, Excel Sales/Orders/Views/Conversion "
                  f"{'= monthly extract' if full else '= direct DB query for those days'}; column + header labels; no undefined/NaN/Infinity",
                  n_all == len(pool) and not bad and lab_ok and sub_ok and not re.search(r"undefined|NaN|Infinity", txt), (n_all, len(pool), lab_ok, sub_ok, bad[:6]))
        # additivity and full-month reconciliation
        add_bad = []
        for k, v in excel[(1, 30)].items():
            for p in ("LY", "TY"):
                for i in range(3):
                    f, x, y = v[p][i], excel[(1, 15)][k][p][i], excel[(16, 30)][k][p][i]
                    if all(isinstance(z, (int, float)) for z in (f, x, y)) and abs((x + y) - f) > (0.0151 if i == 0 else 0.5): add_bad.append((k, p, i, f, x, y))
        check("1–15 Sep + 16–30 Sep = 1–30 Sep for every eBay ID (Sales, Orders, Views, LY and TY)", not add_bad, add_bad[:5])
        if LAST > 30:
            add_bad = []
            for k, v in excel[(1, LAST)].items():
                for p in ("LY", "TY"):
                    for i in range(2):
                        f, x, y = v[p][i], excel[(1, 30)][k][p][i], excel[(31, LAST)][k][p][i]
                        if abs((x + y) - f) > (0.0151 if i == 0 else 0.5): add_bad.append((k, p, i, f, x, y))
            check(f"1–30 Sep + {lbl(31, LAST, 2026)} = {lbl(1, LAST, 2026)} for every eBay ID (Sales, Orders, LY and TY)", not add_bad, add_bad[:5])
        mon = {p: [round(sum(d[p]["s"] for d in pool), 2), sum(d[p]["o"] for d in pool)] for p in ("LY", "TY")}
        f30 = figures["1–30 Sep 2026"]
        check("1–30 Sep totals reconcile exactly with the September extract (dataset.json)",
              all(abs(f30[f"{p} Sales"] - mon[p][0]) < 0.05 and f30[f"{p} Orders"] == mon[p][1] for p in ("LY", "TY")), (mon, f30))
        db30 = {p: [round(sum(v[0] for (mp, i, pp), v in src[(1, 30)].items() if pp == p and (i, mp) in by_key), 2)] for p in ("LY", "TY")}
        check("1–30 Sep totals also equal a direct DB query for the whole month", all(abs(db30[p][0] - f30[f"{p} Sales"]) < 0.05 for p in ("LY", "TY")), (db30, f30))

        # ---- filters with a partial range: PH, category, account, segment, search, pagination ----
        set_range(1, 15); bad = []
        cat = next(c for c in sorted({d["cat"] for d in pool}) if c and sum(1 for d in pool if d["cat"] == c and d["ph"]) > 60)
        for name, c in [("PH utharsika", {"ph": "utharsika"}), (f"Category {cat}", {"cat": cat}), ("Account Sunsone", {"acct": "Sunsone"}),
                        ("Segment A", {"seg": "A – YoY Winner"}), ("search 3179", {"q": "3179"}), ("PH+Account", {"ph": "utharsika", "acct": "Ledsone"})]:
            set_filters(**c); sel = filt(**c)
            if main_n() != len(sel): bad.append((name, "rows", len(sel), main_n())); continue
            fname, wb = download(); b_, _ = verify_excel(wb, fname, sel, (1, 15)); bad += [(name,) + tuple(x) for x in b_]
        check("1–15 Sep with PH / Category / Account / Segment / search filters: rows and Excel values match", not bad, bad[:5])
        set_filters(); pg.select_option("#pgSize", "25"); pg.wait_for_timeout(300); pg.click("#pgNext"); pg.wait_for_timeout(300)
        p2 = pg.inner_text("#pgInfo"); ids2 = pg.eval_on_selector_all("#t tbody tr[data-id]", "r=>r.length")
        check("pagination works with a range selected (page 2, 25 rows)", p2.startswith("Page 2 of") and ids2 == 25 and "1–15 Sep 2026" in pg.inner_text("#count"), (p2, ids2))
        pg.select_option("#pgSize", "50"); pg.wait_for_timeout(200)

        # ---- invalid ranges are blocked ----
        before = pg.eval_on_selector_all("#t tbody tr[data-id] td[data-k=tys]", "e=>e.map(x=>x.innerText)")
        inv = {}
        # one input changed at a time from the valid 1–15 Sep range; each must be rejected without touching the table
        for name, sel_, val in [("From after To", "#sdFrom", "2026-09-20"), ("after the latest data date", "#sdTo", iso(LAST + 1)),
                                ("before 1 Sep", "#sdFrom", "2026-08-25"), ("empty", "#sdFrom", "")]:
            pg.fill("#sdFrom", "2026-09-01"); pg.fill("#sdTo", "2026-09-15"); pg.wait_for_timeout(300)
            pg.fill(sel_, val); pg.wait_for_timeout(400)
            after = pg.eval_on_selector_all("#t tbody tr[data-id] td[data-k=tys]", "e=>e.map(x=>x.innerText)")
            inv[name] = {"msg": pg.inner_text("#rMsg"), "err": "err" in (pg.get_attribute("#rMsg", "class") or ""), "dl_disabled": pg.is_disabled("#fDl"),
                         "table_unchanged": after == before, "still": "1–15 Sep 2026" in pg.inner_text("#count")}
        pg.fill("#sdTo", "2026-09-20"); pg.fill("#sdFrom", "2026-09-20"); pg.wait_for_timeout(400)
        pg.screenshot(path=str(REP / "date_range_single_day.png"))
        recovered = not pg.is_disabled("#fDl") and "20 Sep 2026" in pg.inner_text("#count")
        check("invalid ranges (From > To, after the data end, before 1 Sep, empty) are blocked: error shown, table keeps the last valid range, Download disabled; a valid range recovers",
              all(v["err"] and v["dl_disabled"] and v["table_unchanged"] and v["still"] for v in inv.values()) and recovered, inv)
        set_range(1, 30)
        check("back to 1–30 Sep restores the full-month table", main_n() == len(filt()) and "rng" not in pg.inner_html("#count")
              and all(h.find("|") < 0 for h in pg.eval_on_selector_all("#t thead th", "e=>e.map(x=>x.firstChild.textContent+(x.querySelector('.rh')?'|'+x.querySelector('.rh').textContent:''))")))
        pg.fill("#sdTo", "2026-09-15"); pg.wait_for_timeout(500); pg.screenshot(path=str(REP / "date_range_1_15.png")); set_range(1, 30)
        check("no JS / console errors", not errors, errors[:5])
        br.close()
    ok = all(r["ok"] for r in results)
    out = {"overall": "PASS" if ok else "FAIL", "html": str(HTML), "db": "order_management_copy (DATABASE_URL, read-only)",
           "rule": "TY = selected days of 2026 (1 Sep to the latest order date); LY = the same calendar days of 2025; 1-30 Sep = monthly extract; ads and segment full September",
           "data_end": END, "traffic_end": TRF_END,
           "figures_all_ebay_ids": figures, "checks": results}
    json.dump(out, open(REP / "main_daily_validation.json", "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    json.dump(out, open(REP / f"date_range_validation_{datetime.date.today().isoformat()}.json", "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    print(json.dumps(figures, indent=1, ensure_ascii=False))
    print(f"\n{sum(r['ok'] for r in results)}/{len(results)} passed")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
